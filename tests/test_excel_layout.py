"""Disposició de l'Excel del calendari: totes les setmanes han de fer la
MATEIXA mida i el full ha de quadrar amb un A4 apaïsat."""

import openpyxl
import pandas as pd
import pytest
from openpyxl.utils import range_boundaries

from src.services.excel_export import export_schedule_to_excel

# Àrea útil d'un A4 apaïsat amb els marges de l'exportador (polzades→pt).
PAGE_W, PAGE_H = 11.1 * 72, 7.5 * 72

MACHINES = ["ECO_A", "MX_A", "RM_B"]
# 2026-10-05 és dilluns; tres setmanes seguides.
WEEKS = [
    ["2026-10-05", "2026-10-06", "2026-10-07", "2026-10-08", "2026-10-09"],
    ["2026-10-12", "2026-10-13", "2026-10-14", "2026-10-15", "2026-10-16"],
    ["2026-10-19", "2026-10-20", "2026-10-21", "2026-10-22", "2026-10-23"],
]


def _schedule(with_review_on_weeks=()):
    rows = []
    for w_idx, days in enumerate(WEEKS):
        for day in days:
            for i, slot in enumerate(MACHINES):
                rows.append({
                    "day": day, "franja": "MATI", "slot_id": slot,
                    "professional": "P%d" % (i + 1),
                    "presentiality": "PRESENCIAL", "work_mode": "NORMAL",
                    "is_flipped": 0,
                })
            if w_idx in with_review_on_weeks:
                rows.append({
                    "day": day, "franja": "MATI", "slot_id": "REV_A",
                    "professional": "P1", "presentiality": "NO_PRESENCIAL",
                    "work_mode": "NORMAL", "is_flipped": 0,
                })
    return pd.DataFrame(rows)


def _band_heights(ws):
    """Alçada de cada banda setmanal (una banda comença a la capçalera
    de dia, que és una cel·la fusionada del tipus «Dl 5»)."""
    starts = set()
    for mr in ws.merged_cells.ranges:
        c0, r0, _c1, _r1 = range_boundaries(str(mr))
        v = ws.cell(row=r0, column=c0).value
        if isinstance(v, str) and v[:2] in ("Dl", "Dm", "Dc", "Dj", "Dv"):
            starts.add(r0)
    starts = sorted(starts) + [ws.max_row + 2]
    return [
        sum((ws.row_dimensions[r].height or 15)
            for r in range(starts[i], starts[i + 1] - 1))
        for i in range(len(starts) - 1)
    ]


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    """Directori de treball amb un catàleg on REV_A és REVISIÓ. Cal: el
    caràcter de revisió surt del catàleg (columna `review`), mai del nom,
    i l'exportador el llegeix del directori actual."""
    monkeypatch.chdir(tmp_path)
    (tmp_path / "data").mkdir()
    pd.DataFrame([
        {"slot_id": s, "weekday": 1, "weekend": 0, "linked_to": "",
         "doubled": 0, "review": int(s == "REV_A"), "always_presential": 0,
         "area": "ZONA_A", "metric_family": "", "assignee": "", "notes": ""}
        for s in MACHINES + ["REV_A"]
    ]).to_csv(tmp_path / "data" / "slot_catalog.csv",
              index=False, encoding="utf-8-sig")
    return tmp_path


def _export(tmp_path, df, months=(10,)):
    csv = tmp_path / "sched.csv"
    df.to_csv(csv, index=False)
    out = tmp_path / "cal.xlsx"
    export_schedule_to_excel(csv, out, selected_months=list(months), year=2026)
    return openpyxl.load_workbook(out)


def _page_fill(ws):
    """(% d'ample, % d'alt) de la pàgina que ocupa el full en imprimir-se
    amb «ajustar a una pàgina»."""
    w = sum(ws.column_dimensions[c].width or 0
            for c in ws.column_dimensions) * 5.25
    h = sum((ws.row_dimensions[r].height or 15)
            for r in range(1, ws.max_row + 1))
    scale = min(PAGE_W / w, PAGE_H / h)
    return w * scale / PAGE_W * 100, h * scale / PAGE_H * 100


class TestUniformWeekBands:
    def test_all_weeks_same_height_without_reviews(self, workspace):
        ws = _export(workspace, _schedule())["Calendari"]
        alts = _band_heights(ws)
        assert len(alts) == 3
        assert len(set(round(a, 1) for a in alts)) == 1

    def test_all_weeks_same_height_when_review_only_in_some(self, workspace):
        # REGRESSIÓ: les files de revisió es calculaven per setmana, així
        # que una setmana sense revisió quedava ~20% més curta.
        ws = _export(workspace, _schedule(with_review_on_weeks=(0, 2)))["Calendari"]
        alts = _band_heights(ws)
        assert len(alts) == 3
        assert len(set(round(a, 1) for a in alts)) == 1, (
            "setmanes de mides diferents: %s" % alts
        )

    def test_review_row_present_in_every_band(self, workspace):
        ws = _export(workspace, _schedule(with_review_on_weeks=(1,)))["Calendari"]
        etiquetes = [
            str(ws.cell(row=r, column=1).value or "")
            for r in range(1, ws.max_row + 1)
        ]
        assert sum(1 for e in etiquetes if e.startswith("Rev.")) == 3


class TestA4Fit:
    @pytest.mark.parametrize("n_weeks", [1, 2, 3])
    def test_sheet_fills_the_page(self, workspace, n_weeks):
        df = _schedule()
        keep = {d for days in WEEKS[:n_weeks] for d in days}
        ws = _export(workspace, df[df["day"].isin(keep)])["Calendari"]
        ample, alt = _page_fill(ws)
        # Abans del segon ajust d'amplada, un mes d'una sola setmana
        # ocupava ~38% de l'alçada del full.
        assert min(ample, alt) >= 70, (
            "full desquadrat amb %d setmana(es): %.0f%% x %.0f%%"
            % (n_weeks, ample, alt)
        )

    def test_one_sheet_per_month_each_fitted(self, workspace):
        df = _schedule()
        nov = df.copy()
        nov["day"] = nov["day"].str.replace("2026-10-", "2026-11-", regex=False)
        wb = _export(workspace, pd.concat([df, nov]), months=(10, 11))
        assert len(wb.sheetnames) == 2
        for name in wb.sheetnames:
            ample, alt = _page_fill(wb[name])
            assert min(ample, alt) >= 70
